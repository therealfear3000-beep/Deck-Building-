#!/usr/bin/env python3
"""Structural/heuristic PPTX audit. This is NOT a substitute for rendering.

Usage: python audit_pptx.py deck.pptx --mode read --json audit.json
Exit 1: structural errors. Exit 0 may still contain warnings requiring review.
Requires python-pptx; Pillow/fontconfig improve optional text-fit estimates.
"""
import argparse
from collections import Counter
from functools import lru_cache
import json
from pathlib import Path
import posixpath
import re
import subprocess
import sys
from urllib.parse import unquote
import zipfile
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "r": "http://schemas.openxmlformats.org/package/2006/relationships"}


@lru_cache(maxsize=128)
def load_font(family, points):
    try:
        from PIL import ImageFont
        result = subprocess.run(["fc-match", "-f", "%{file}", family],
                                capture_output=True, text=True, check=True, timeout=5)
        return ImageFont.truetype(result.stdout.strip(), max(1, round(points*96/72)))
    except (ImportError, OSError, ValueError, subprocess.SubprocessError):
        return None


def style(p, run=None):
    size = ((run.font.size if run is not None else None) or p.font.size)
    name = ((run.font.name if run is not None else None) or p.font.name or "Arial")
    return (size.pt if size is not None else None), name


def text_fit(tf, width, height):
    """Conservative word wrapping estimate; engine differences make this a warning."""
    available_w = (width-tf.margin_left-tf.margin_right)/EMU*96
    available_h = (height-tf.margin_top-tf.margin_bottom)/EMU*96
    if available_w <= 0 or available_h <= 0:
        return True, "non-positive inner text area"
    total_h = 0
    for p in tf.paragraphs:
        if not p.text:
            continue
        sizes = [style(p, r)[0] for r in p.runs if r.text] or [style(p)[0]]
        if any(x is None for x in sizes):
            return False, "unresolved inherited font; inspect render"
        size = max(sizes)
        family = style(p, p.runs[0] if p.runs else None)[1]
        first_run = p.runs[0] if p.runs else None
        bold = first_run.font.bold if first_run is not None else None
        italic = first_run.font.italic if first_run is not None else None
        bold = p.font.bold if bold is None else bold
        italic = p.font.italic if italic is None else italic
        if bold or italic:
            family += ":style=" + " ".join(x for x, present in [("Bold", bold), ("Italic", italic)] if present)
        font = load_font(family, size)
        if font is None:
            return False, "font metrics unavailable; inspect render"
        lines = 0
        for raw in p.text.split("\v"):
            current = ""
            for word in raw.split():
                if font.getlength(word) > available_w*1.04:
                    return True, "unbroken text may exceed box width"
                candidate = current+" "+word if current else word
                if font.getlength(candidate) > available_w and current:
                    lines += 1
                    current = word
                else:
                    current = candidate
            lines += 1
        spacing = p.line_spacing
        if isinstance(spacing, float):
            line_h = size*96/72*spacing
        elif spacing is not None:
            line_h = spacing.pt*96/72
        else:
            line_h = size*96/72*1.15
        before = p.space_before.pt if p.space_before is not None else 0
        after = p.space_after.pt if p.space_after is not None else 0
        total_h += lines*line_h+(before+after)*96/72
    return total_h > available_h*1.12, f"estimated text height {total_h:.0f}px vs {available_h:.0f}px"


def leaf_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from leaf_shapes(shape.shapes)
        else:
            yield shape


def audit(path, mode="read"):
    report = {"file": str(path), "mode": mode, "errors": [], "warnings": [],
              "slides": [], "limitations": [
                  "Text fit is heuristic; visual rendering is mandatory.",
                  "Semantic accuracy, financial reconciliation, contrast, reading order, and accidental overlap require human/model review.",
                  "Bounds check applies to top-level objects; inspect children of groups visually.",
                  "Native object presence does not prove that every visible label is editable."]}

    def add(level, slide, obj, message):
        report[level].append({"slide": slide, "object": obj, "message": message})

    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            add("errors", None, bad, "ZIP integrity failure")
        names = set(z.namelist())
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    root = etree.fromstring(z.read(name))
                except etree.XMLSyntaxError as e:
                    add("errors", None, name, f"Invalid XML: {e}")
                    continue
                if name.endswith(".rels"):
                    base = "" if name == "_rels/.rels" else name.rsplit("/_rels/", 1)[0]
                    for rel in root:
                        if rel.get("TargetMode") == "External":
                            continue
                        target = unquote(rel.get("Target", "")).split("#", 1)[0]
                        resolved = (target.lstrip("/") if target.startswith("/") else
                                    posixpath.normpath(posixpath.join(base, target)))
                        if resolved not in names:
                            add("errors", None, name, f"Missing internal relationship target: {resolved}")
        report["native_chart_files"] = sum(bool(re.fullmatch(r"ppt/charts/chart\d+\.xml", n)) for n in names)
        report["embedded_workbooks"] = sum(n.startswith("ppt/embeddings/") and n.endswith(".xlsx") for n in names)
        for name in sorted(n for n in names if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)):
            root = etree.fromstring(z.read(name))
            if not root.findall(".//c:externalData", NS):
                add("warnings", None, name, "Chart has no linked embedded data relationship; verify Edit Data works")
    prs = Presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    for idx, slide in enumerate(prs.slides, 1):
        counts = Counter()
        body_words = 0
        for shape in slide.shapes:
            if (shape.left < -EMU*.02 or shape.top < -EMU*.02 or
                    shape.left+shape.width > sw+EMU*.02 or shape.top+shape.height > sh+EMU*.02):
                add("errors", idx, shape.name, "Object extends beyond slide bounds")
        for shape in leaf_shapes(slide.shapes):
            role = shape.name.split(":")[0].lower()
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counts["pictures"] += 1
                if shape.width*shape.height > sw*sh*.85:
                    add("warnings", idx, shape.name, "Near-full-slide picture: verify this is not a flattened content slide")
            elif shape.has_chart:
                counts["charts"] += 1
                chart_text = shape.chart._chartSpace.xpath(".//a:t | .//c:strCache/c:pt/c:v")
                body_words += sum(len(e.text.split()) for e in chart_text if e.text)
            elif shape.has_table:
                counts["tables"] += 1
            elif shape.has_text_frame and shape.text.strip():
                counts["text_objects"] += 1
            else:
                counts["other_native_shapes"] += 1
            frames = []
            if shape.has_text_frame and shape.text.strip():
                frames.append((shape.text_frame, shape.width, shape.height, role, shape.name))
            if shape.has_table:
                for ri, row in enumerate(shape.table.rows):
                    for ci, cell in enumerate(row.cells):
                        if cell.is_spanned:
                            continue
                        width = sum(shape.table.columns[k].width for k in range(ci, ci+cell.span_width))
                        height = sum(shape.table.rows[k].height for k in range(ri, ri+cell.span_height))
                        frames.append((cell.text_frame, width, height, "table", f"{shape.name}[{ri},{ci}]"))
            for tf, width, height, kind, name in frames:
                if kind not in {"footer", "nav", "title"}:
                    body_words += len(re.findall(r"\S+", tf.text))
                threshold = (8 if kind in {"footer", "nav"} else
                             14 if kind in {"label", "table", "takeaway"} else
                             20 if mode == "live" else 16)
                sizes = [style(p, r)[0] for p in tf.paragraphs for r in p.runs if r.text.strip()]
                known = [s for s in sizes if s is not None]
                if known and min(known) < threshold:
                    add("warnings", idx, name, f"Font {min(known):g}pt below role heuristic {threshold}pt")
                if any(s is None for s in sizes):
                    add("warnings", idx, name, "Inherited font size unresolved; check theme and rendering")
                overflow, reason = text_fit(tf, width, height)
                if overflow:
                    add("warnings", idx, name, f"Possible text overflow: {reason}")
        limit = 120 if mode == "live" else 220
        if body_words > limit:
            add("warnings", idx, "slide", f"Approximate visible body words {body_words} exceed {limit}; review density")
        if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
            add("warnings", idx, "notes", "No substantive notes found; add source/assumption detail where relevant")
        if counts["pictures"] and not counts["text_objects"] and not counts["tables"] and not counts["charts"]:
            add("warnings", idx, "slide", "Image-only content: investigate editability")
        report["slides"].append({"slide": idx, "body_words_approx": body_words, "objects": dict(counts)})
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--mode", choices=["read", "live"], default="read")
    parser.add_argument("--json", type=Path)
    args = parser.parse_args()
    try:
        result = audit(args.pptx, args.mode)
    except Exception as e:
        print(f"Audit failed: {type(e).__name__}: {e}", file=sys.stderr)
        return 2
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))
    return 1 if result["errors"] else 0


if __name__ == "__main__":
    sys.exit(main())

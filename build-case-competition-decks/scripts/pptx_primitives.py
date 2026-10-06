"""Optional native PowerPoint primitives; author layouts, do not fill a fixed template.

Requires python-pptx. Coordinates are inches, font sizes points, colours hex RGB.
All business text, tables, connectors, diagrams and charts stay editable.
"""
from dataclasses import dataclass
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class Theme:
    font: str = "Arial"
    ink: str = "20243A"
    primary: str = "542785"
    accent: str = "007C83"
    surface: str = "F3F1F7"
    muted: str = "596273"
    line: str = "D9DCE5"
    white: str = "FFFFFF"
    body_pt: float = 17


def rgb(value):
    return RGBColor.from_string(value.lstrip("#").upper())


class Deck:
    def __init__(self, theme=None, width=13.333333, height=7.5):
        self.theme = theme or Theme()
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(width), Inches(height)
        self.width, self.height = width, height
        self.prs.core_properties.title = "Case competition presentation"
        self.prs.core_properties.author = ""

    def text(self, slide, value, x, y, w, h, size=None, bold=False,
             color=None, role="body", align="left", margin=0.02):
        if min(w, h) <= 0:
            raise ValueError("Text box dimensions must be positive")
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = f"{role}:{str(value)[:55]}"
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(margin)
        tf.margin_top = tf.margin_bottom = Inches(margin)
        lines = str(value).split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = self.theme.font
            p.font.size = Pt(size if size is not None else self.theme.body_pt)
            p.font.bold = bold
            p.font.color.rgb = rgb(color or self.theme.ink)
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                           "right": PP_ALIGN.RIGHT}[align]
            p.space_before = Pt(0)
            p.space_after = Pt(4 if i < len(lines)-1 else 0)
            p.line_spacing = 1.08
        return shape

    def box(self, slide, x, y, w, h, fill=None, line=None,
            kind=MSO_SHAPE.RECTANGLE, name="panel"):
        shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
        else:
            shape.fill.background()
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        shape._element.spPr.append(OxmlElement("a:effectLst"))
        for ref in shape._element.xpath("./p:style/a:effectRef"):
            ref.set("idx", "0")
        return shape

    def connector(self, slide, x1, y1, x2, y2, color=None, arrow=True,
                  width=1.5, begin_shape=None, end_shape=None):
        sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                       Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        sh.name = "diagram:connector"
        sh.line.color.rgb = rgb(color or self.theme.accent)
        sh.line.width = Pt(width)
        if begin_shape is not None:
            sh.begin_connect(begin_shape, 3)
        if end_shape is not None:
            sh.end_connect(end_shape, 1)
        if arrow:
            end = OxmlElement("a:tailEnd")
            end.set("type", "triangle")
            sh._element.spPr.get_or_add_ln().append(end)
        sh._element.spPr.append(OxmlElement("a:effectLst"))
        for ref in sh._element.xpath("./p:style/a:effectRef"):
            ref.set("idx", "0")
        return sh

    def base(self, title, section="", source="", takeaway=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(self.theme.white)
        self.box(slide, 0, 0, self.width, .09, fill=self.theme.primary, name="decor:top-rule")
        self.text(slide, section.upper(), .45, .24, self.width-1, .24,
                  size=11, bold=True, color=self.theme.accent, role="nav")
        self.text(slide, title, .45, .59, self.width-.90, 1.02,
                  size=28, bold=True, role="title")
        self.connector(slide, .45, 1.72, self.width-.45, 1.72,
                       color=self.theme.line, arrow=False, width=.8)
        if takeaway:
            self.box(slide, .45, self.height-.92, self.width-.9, .43,
                     fill=self.theme.surface, name="panel:takeaway")
            self.text(slide, takeaway, .59, self.height-.89, self.width-1.2, .36,
                      size=15, bold=True, color=self.theme.primary, role="takeaway")
        self.text(slide, source, .45, self.height-.30, self.width-1.35, .21,
                  size=9, color=self.theme.muted, role="footer")
        self.text(slide, str(len(self.prs.slides)), self.width-.78, self.height-.32,
                  .33, .24, size=10, color=self.theme.muted, role="footer", align="right")
        return slide

    def panel(self, slide, heading, body, x, y, w, h, accent=None):
        self.box(slide, x, y, w, h, fill=self.theme.surface, name=f"panel:{heading}")
        self.box(slide, x, y, .045, h, fill=accent or self.theme.primary, name="decor:panel-rule")
        self.text(slide, heading, x+.18, y+.16, w-.36, .50,
                  size=19, bold=True, color=accent or self.theme.primary, role="heading")
        self.text(slide, body, x+.18, y+.76, w-.36, h-.92)

    def metric(self, slide, value, label, x, y, w=2.8, note=""):
        self.text(slide, value, x, y, w, .70, size=36, bold=True,
                  color=self.theme.primary, role="metric")
        self.text(slide, label, x, y+.80, w, .64, size=17, bold=True, role="label")
        if note:
            self.text(slide, note, x, y+1.50, w, .64, size=15, role="label")

    def flow(self, slide, steps, x, y, w, h=2.5):
        """Editable process with 2–5 (heading, explanation) tuples."""
        if not 2 <= len(steps) <= 5:
            raise ValueError("Use 2–5 steps; split larger flows")
        gap = .32
        card_w = (w-gap*(len(steps)-1))/len(steps)
        shapes = []
        for i, (head, body) in enumerate(steps):
            sx = x+i*(card_w+gap)
            sh = self.box(slide, sx, y, card_w, h, fill=self.theme.surface,
                          name=f"diagram:step-{i+1}-{head}")
            shapes.append(sh)
            self.text(slide, f"{i+1:02d}", sx+.15, y+.14, .50, .35,
                      size=17, bold=True, color=self.theme.accent, role="label")
            self.text(slide, head, sx+.15, y+.60, card_w-.30, .68,
                      size=18, bold=True, role="heading")
            self.text(slide, body, sx+.15, y+1.40, card_w-.30, h-1.55, size=16)
        for left, right in zip(shapes, shapes[1:]):
            self.connector(slide, left.left/914400+card_w, y+h/2,
                           right.left/914400, y+h/2, begin_shape=left, end_shape=right)
        return shapes

    def table(self, slide, headers, rows, x, y, w, h, widths=None, size=16):
        if not headers or any(len(r) != len(headers) for r in rows):
            raise ValueError("Every table row must match the headers")
        shape = slide.shapes.add_table(len(rows)+1, len(headers),
                                      Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = "table:editable"
        table = shape.table
        if widths is not None:
            if len(widths) != len(headers) or min(widths) <= 0:
                raise ValueError("Column weights must match headers and be positive")
            for col, weight in zip(table.columns, widths):
                col.width = Inches(w*weight/sum(widths))
        for ri, values in enumerate([headers]+list(rows)):
            for ci, value in enumerate(values):
                cell = table.cell(ri, ci)
                cell.text = str(value)
                cell.margin_left = cell.margin_right = Inches(.10)
                cell.margin_top = cell.margin_bottom = Inches(.07)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(self.theme.primary if ri == 0 else
                                                (self.theme.surface if ri%2 else self.theme.white))
                for p in cell.text_frame.paragraphs:
                    p.font.name = self.theme.font
                    p.font.size = Pt(size)
                    p.font.bold = (ri == 0)
                    p.font.color.rgb = rgb(self.theme.white if ri == 0 else self.theme.ink)
                    p.space_before = p.space_after = Pt(0)
        return shape

    def bar_chart(self, slide, categories, series, x, y, w, h,
                  number_format="0", horizontal=True):
        """Series is an ordered mapping name -> numeric sequence; native embedded workbook."""
        if not categories or not series:
            raise ValueError("Provide categories and at least one series")
        data = CategoryChartData()
        data.categories = categories
        for name, values in series.items():
            if len(values) != len(categories):
                raise ValueError("Chart categories and series lengths differ")
            data.add_series(name, values)
        kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        shape = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
        shape.name = "chart:editable-with-workbook"
        chart = shape.chart
        chart.has_title = False
        chart.has_legend = len(series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(14)
        for axis in (chart.category_axis, chart.value_axis):
            axis.tick_labels.font.name = self.theme.font
            axis.tick_labels.font.size = Pt(15)
            axis.tick_labels.font.color.rgb = rgb(self.theme.ink)
        values = [v for seq in series.values() for v in seq if v is not None]
        if values and min(values) >= 0:
            chart.value_axis.minimum_scale = 0
        chart.value_axis.tick_labels.number_format = number_format
        chart.category_axis.has_major_gridlines = False
        chart.plots[0].has_data_labels = True
        labels = chart.plots[0].data_labels
        labels.font.size = Pt(15)
        labels.font.name = self.theme.font
        labels.number_format = number_format
        labels.position = XL_LABEL_POSITION.OUTSIDE_END
        colours = [self.theme.primary, self.theme.accent, "78869B", "B96A09"]
        for idx, s in enumerate(chart.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(colours[idx % len(colours)])
            s.format.line.fill.background()
        return shape

    def icon(self, slide, kind, x, y, size=.45, color=None):
        """Small editable cues; returns native shapes. Use matching labels nearby."""
        c = color or self.theme.accent
        if kind == "person":
            return [self.box(slide, x+size*.32, y, size*.36, size*.36, fill=c,
                             kind=MSO_SHAPE.OVAL, name="icon:person-head"),
                    self.box(slide, x+size*.14, y+size*.42, size*.72, size*.55, fill=c,
                             kind=MSO_SHAPE.ROUNDED_RECTANGLE, name="icon:person-body")]
        if kind == "gate":
            return [self.box(slide, x, y, size, size, fill=c, kind=MSO_SHAPE.DIAMOND,
                             name="icon:decision-gate")]
        if kind == "target":
            return [self.box(slide, x, y, size, size, line=c, kind=MSO_SHAPE.OVAL,
                             name="icon:target-outer"),
                    self.box(slide, x+size*.30, y+size*.30, size*.40, size*.40, fill=c,
                             kind=MSO_SHAPE.OVAL, name="icon:target-centre")]
        raise ValueError("Supported icons: person, gate, target")

    @staticmethod
    def notes(slide, text):
        slide.notes_slide.notes_text_frame.text = text

    def save(self, path):
        self.prs.save(str(path))
